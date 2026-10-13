"""Build slide 03 (odor-information gap) as an editable PPTX."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import copy

PX = 1.0 / 96.0
def I(px): return Inches(px * PX)

C = dict(purple="6B3F99", teal="1F8A8A", pink="D94F8A", red="C2410C", green="15803D",
         ink="2B2B2B", mute="6E6E73", t_teal="E9F4F4", t_purple="F3F0F7",
         t_pink="FCEEF4", gray="EDEDED", white="FFFFFF", dash="BFBFC8", mtas="E8412B")
def rgb(k): return RGBColor.from_string(C[k])
FONT = "Verdana"

prs = Presentation()
prs.slide_width, prs.slide_height = I(1280), I(720)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes

def textbox(x, y, w, h, runs, size=14, color="ink", bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, spacing=1.0, italic=False):
    tb = S.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [[(runs, bold)]]
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for text, b in line:
            r = p.add_run(); r.text = text
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(size), b, italic
            f.color.rgb = rgb(color)
    return tb

def box(x, y, w, h, fill, radius=0.06, line=None, lw=2, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = S.add_shape(shape, I(x), I(y), I(w), I(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh

def dashed(sh):
    ln = sh.line._get_or_add_ln()
    d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'}); ln.append(d)

# ---------- header ----------
textbox(72, 42, 1000, 46, "How do we capture odor differences?", size=25, color="purple", bold=True)
textbox(74, 90, 1000, 24, "Machines copy the nose — so they inherit its blind spot.",
        size=11, color="mute")

# ---------- columns ----------
COLS = [
    dict(x=72,  tint="t_teal",   tag="Human nose",          tagc="teal",
         ph=["Cell 1999 figure here", "odorants × receptors matrix", "(crop to ~5 rows)"],
         cite="Malnic et al., Cell 96, 713 (1999)",
         cap=[[("One odorant hits ", False), ("many receptors", True), (".", False)],
              [("Identity is the pattern.", True)]],
         chips=[("✓ pattern", "green"), ("✗ chemistry", "red")], callout=None, art=None),
    dict(x=457, tint="t_purple", tag="Machines today",      tagc="purple",
         ph=["Adv. Mater. 2024 figure here", "panel b only: sensor array", "→ pattern recognition"],
         cite="Adv. Mater. 36, 2402287 (2024)",
         cap=[[("Arrays rebuild ", False), ("the same pattern", True), (",", False)],
              [("with ", False), ("less", True), (" chemical detail.", False)]],
         chips=[("✓ pattern", "green"), ("✗ chemistry", "red")],
         callout="pattern only — which molecule?", art=None),
    dict(x=842, tint="t_pink",   tag="Raman — this work", tagc="pink",
         ph=None, cite="",
         cap=[[("Hundreds of narrow bands at once.", False)],
              [("Pattern and chemistry together.", True)]],
         chips=[("✓ pattern", "green"), ("✓ chemistry", "green")],
         callout=None, art="raman-art.png"),
]
CW, CTOP, CH = 365, 130, 408

for i, c in enumerate(COLS):
    outline = "pink" if i == 2 else None
    box(c["x"], CTOP, CW, CH, c["tint"], radius=0.05, line=outline, lw=2.5)

    # tag
    tagw = 150 if i == 0 else (168 if i == 1 else 196)
    tx = c["x"] + (CW - tagw) / 2
    t = box(tx, CTOP + 16, tagw, 30, c["tagc"], radius=0.28)
    tf = t.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = c["tag"]
    r.font.name, r.font.size, r.font.bold = FONT, Pt(10.5), True
    r.font.color.rgb = rgb("white")

    # figure slot
    sx, sy, sw, sh_ = c["x"] + 18, CTOP + 57, CW - 36, 184
    slot = box(sx, sy, sw, sh_, "white", radius=0.05,
               line=("dash" if c["ph"] else None), lw=1.5)
    if c["ph"]:
        dashed(slot)
        textbox(sx, sy + 58, sw, 70,
                [[(c["ph"][0], True)], [(c["ph"][1], False)], [(c["ph"][2], False)]],
                size=9, color="mute", align=PP_ALIGN.CENTER, spacing=1.35)
    if c["art"]:
        S.add_picture(c["art"], I(sx + 6), I(sy + 10), width=I(sw - 12))
    if c["callout"]:
        co = box(sx + sw - 232, sy + sh_ - 34, 224, 24, "white", radius=0.4, line="red", lw=1)
        tfc = co.text_frame; tfc.vertical_anchor = MSO_ANCHOR.MIDDLE
        pc = tfc.paragraphs[0]; pc.alignment = PP_ALIGN.CENTER
        rc = pc.add_run(); rc.text = c["callout"]
        rc.font.name, rc.font.size, rc.font.bold = FONT, Pt(8.5), True
        rc.font.color.rgb = rgb("red")

    # citation
    if c["cite"]:
        textbox(c["x"] + 10, CTOP + 247, CW - 20, 16, c["cite"],
                size=8, color="mute", align=PP_ALIGN.CENTER, italic=True)

    # caption
    textbox(c["x"] + 14, CTOP + 272, CW - 28, 50, c["cap"],
            size=10.5, color="ink", align=PP_ALIGN.CENTER, spacing=1.35)

    # chips
    cw1, cw2, gap = 112, 124, 10
    total = cw1 + cw2 + gap
    bx = c["x"] + (CW - total) / 2
    for j, (label, col) in enumerate(c["chips"]):
        w = cw1 if j == 0 else cw2
        ch_ = box(bx, CTOP + 338, w, 28, "white", radius=0.45)
        tfh = ch_.text_frame; tfh.vertical_anchor = MSO_ANCHOR.MIDDLE
        ph_ = tfh.paragraphs[0]; ph_.alignment = PP_ALIGN.CENTER
        rh = ph_.add_run(); rh.text = label
        rh.font.name, rh.font.size, rh.font.bold = FONT, Pt(10), True
        rh.font.color.rgb = rgb(col)
        bx += w + gap

# ---------- verdict ----------
box(72, 558, 1136, 72, "gray", radius=0.14)
textbox(100, 574, 1080, 24,
        "GC–MS has the chemistry but not the pattern; the nose and the e-nose have the pattern but not the chemistry.",
        size=11, color="mute", align=PP_ALIGN.CENTER)
textbox(100, 600, 1080, 24,
        [[("→  ", False), ("Read the Raman spectrum as a fingerprint map.", True)]],
        size=12, color="pink", align=PP_ALIGN.CENTER)

# ---------- footer ----------
textbox(76, 646, 320, 40, "μTAS 2026", size=21, color="mtas", bold=True)
textbox(1150, 658, 100, 20, "03", size=9, color="mute", align=PP_ALIGN.RIGHT)

band = S.add_shape(MSO_SHAPE.RECTANGLE, I(0), I(696), I(1280), I(24))
band.line.fill.background(); band.shadow.inherit = False
band.fill.gradient()
stops = band.fill.gradient_stops
stops[0].color.rgb = RGBColor.from_string("1E6B35"); stops[0].position = 0.0
stops[1].color.rgb = RGBColor.from_string("FFFFFF"); stops[1].position = 1.0
band.fill.gradient_angle = 0.0

prs.save("slide03-gap.pptx")
print("saved slide03-gap.pptx")
