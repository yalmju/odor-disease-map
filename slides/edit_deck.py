"""Rebuild slide 4 (the 'how do we capture odor differences' gap slide) inside the real deck."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

SRC, DST = "OpticalNose_5min_v7.pptx", "OpticalNose_5min_v8.pptx"
prs = Presentation(SRC)
slide = prs.slides[3]
S = slide.shapes

C = dict(purple="6B3F99", teal="1F8A8A", amber="C97A1E", pink="D94F8A",
         red="C2410C", green="15803D", ink="2B2B2B", mute="6E6E73",
         t_teal="E9F4F4", t_purple="F3F0F7", t_amber="FDF3E7",
         gray="EDEDED", white="FFFFFF")
def rgb(k): return RGBColor.from_string(C[k])

# ---------- 1. clear the old pasted figure group ----------
for sh in list(S):
    if sh.shape_type == 6:                      # GROUP
        sh._element.getparent().remove(sh._element)

# ---------- 2. title ----------
for sh in S:
    if sh.has_text_frame and "odor difference" in sh.text_frame.text:
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = "How do we capture odor differences?"

def textbox(x, y, w, h, lines, size=12, color="ink", align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, spacing=1.0, italic=False):
    tb = S.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [[(lines, False)]]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        for text, bold in line:
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
            r.font.color.rgb = rgb(color)
    return tb

def box(x, y, w, h, fill, radius=0.08, line=None, lw=1.5):
    sh = S.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh

def pill(x, y, w, h, fill, label, size=13):
    sh = box(x, y, w, h, fill, radius=0.3)
    tf = sh.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = rgb("white")
    return sh

# ---------- 3. subtitle ----------
textbox(0.92, 1.12, 11.5, 0.3,
        "Three ways to read a smell today — each one keeps only part of it.",
        size=16, color="mute")

# ---------- 4. columns ----------
L, TOTAL, GAP = 0.90, 11.53, 0.24
CW = (TOTAL - 2 * GAP) / 3
TOP, CH = 1.52, 3.98

COLS = [
    dict(tint="t_teal", accent="teal", tag="Human nose", tagw=1.95,
         img="fig_cell.png",
         cite="Malnic et al., Cell 96, 713 (1999)",
         cap=[[("One odorant hits ", False), ("many receptors", True), (".", False)],
              [("Identity is the pattern.", True)]],
         chips=[("✓  pattern", "green", 1.30), ("✗  chemistry", "red", 1.45)]),
    dict(tint="t_purple", accent="purple", tag="Artificial nose", tagw=2.25,
         img="fig_advmat.png",
         cite="Adv. Mater. 36, 2402287 (2024)",
         cap=[[("Sensor arrays copy ", False), ("the same pattern", True), (",", False)],
              [("with ", False), ("less", True), (" chemical detail.", False)]],
         chips=[("✓  pattern", "green", 1.30), ("✗  chemistry", "red", 1.45)]),
    dict(tint="t_amber", accent="amber", tag="Odor models", tagw=2.00,
         img="fig_pom.png",
         cite="Lee et al., Science 381, 6661 (2023)",
         cap=[[("A learned map predicts odor,", False)],
              [("but the ", False), ("molecule must be known", True), (".", False)]],
         chips=[("✓  odor", "green", 1.15), ("✗  from a sample", "red", 1.80)]),
]

for i, c in enumerate(COLS):
    x = L + i * (CW + GAP)
    box(x, TOP, CW, CH, c["tint"], radius=0.05)

    pill(x + (CW - c["tagw"]) / 2, TOP + 0.12, c["tagw"], 0.34, c["accent"], c["tag"])

    # white figure slot
    sx, sy = x + 0.16, TOP + 0.55
    sw, sh_ = CW - 0.32, 1.80
    box(sx, sy, sw, sh_, "white", radius=0.05)

    # fit the picture inside the slot
    iw, ih = Image.open(c["img"]).size
    ar = iw / ih
    maxw, maxh = sw - 0.18, sh_ - 0.16
    w = maxw; h = w / ar
    if h > maxh: h = maxh; w = h * ar
    S.add_picture(c["img"], Inches(sx + (sw - w) / 2), Inches(sy + (sh_ - h) / 2),
                  width=Inches(w), height=Inches(h))

    textbox(x + 0.10, TOP + 2.42, CW - 0.20, 0.22, c["cite"],
            size=9, color="mute", align=PP_ALIGN.CENTER, italic=True)

    textbox(x + 0.14, TOP + 2.66, CW - 0.28, 0.60, c["cap"],
            size=13, color="ink", align=PP_ALIGN.CENTER, spacing=1.3)

    total_w = sum(w for _, _, w in c["chips"]) + 0.12
    bx = x + (CW - total_w) / 2
    for label, col, w in c["chips"]:
        ch_ = box(bx, TOP + 3.30, w, 0.32, "white", radius=0.45)
        tf = ch_.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = rgb(col)
        bx += w + 0.12

# ---------- 5. bottom bar ----------
bar = box(L, 5.60, TOTAL, 0.88, "gray", radius=0.16)
tf = bar.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = tf.margin_right = Inches(0.25)
p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
r = p1.add_run(); r.text = "A pattern without chemistry, or chemistry without a measurement."
r.font.size = Pt(15); r.font.color.rgb = rgb("mute")
p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run(); r.text = "→  Can we measure a real sample and keep both?"
r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = rgb("pink")

# ---------- 6. fix page numbers on every slide ----------
for n, sl in enumerate(prs.slides, start=1):
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.left > Inches(12):
            para = sh.text_frame.paragraphs[0]
            if para.runs:
                para.runs[0].text = f"{n:02d}"
                for extra in para.runs[1:]:
                    extra.text = ""

prs.save(DST)
print("saved", DST)
